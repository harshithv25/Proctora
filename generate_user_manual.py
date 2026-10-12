import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.abspath("docs")
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "user_manual_screenshots")
DOCX_OUTPUT = os.path.join(DOCS_DIR, "Proctora_User_Manual.docx")
HTML_OUTPUT = os.path.join(DOCS_DIR, "Proctora_User_Manual.html")
PDF_OUTPUT = os.path.join(DOCS_DIR, "Proctora_User_Manual.pdf")

authors = [
    {"name": "Adarsh Bellamane", "roll": "241IT004"},
    {"name": "Harshith Vellapha", "roll": "241IT033"},
    {"name": "Rushi Hiren Patel", "roll": "241IT065"},
]

sections = [
    {
        "flow": "Flow 1: System Landing Page & Public Navigation",
        "description": "This flow covers accessing the platform, viewing system capabilities, and navigating to candidate or administrator authentication.",
        "steps": [
            {
                "title": "Step 1.1: System Landing Page",
                "image": "step_01_landing_page.png",
                "text": "The landing page introduces Proctora's real-time online examination features, including AI-driven gaze tracking, Monaco editor telemetry, and classroom seating deconfliction. Users can click 'Candidate Portal' to take tests or 'Admin Login' to manage examinations."
            }
        ]
    },
    {
        "flow": "Flow 2: User Registration & Authentication",
        "description": "This flow covers candidate account registration and secure credential login with role-based redirection.",
        "steps": [
            {
                "title": "Step 2.1: Candidate Account Registration",
                "image": "step_02_registration_page.png",
                "text": "New candidates register by providing their full name, institutional email address, institutional roll number, and password. Client and server validation verify password strength and email formatting."
            },
            {
                "title": "Step 2.2: Candidate & User Login",
                "image": "step_03_login_page.png",
                "text": "Candidates sign in using their registered email and password. Upon successful verification, candidates are issued secure session cookies and redirected to the Candidate Dashboard."
            },
            {
                "title": "Step 2.3: Administrator Login Interface",
                "image": "step_13_admin_login.png",
                "text": "Faculty and administrators log in through the same secure authentication gateway. Users with administrative privileges are granted access to the Test Studio and Live Invigilation Monitor."
            }
        ]
    },
    {
        "flow": "Flow 3: Candidate Examination Flow",
        "description": "This flow details the complete candidate testing lifecycle from test entry, hardware calibration, answering diverse question formats, to final submission.",
        "steps": [
            {
                "title": "Step 3.1: Candidate Dashboard & Entering Test ID",
                "image": "step_04_candidate_dashboard.png",
                "text": "Upon logging in, candidates view a minimalist portal with an input field labeled 'Paste your exam link or Test ID here:'. Candidates enter the assigned Test ID (e.g., TEST-MS2026) and click 'Proceed to Examination'."
            },
            {
                "title": "Step 3.2: Pre-Flight Hardware Diagnostics & AI Face Calibration",
                "image": "step_05_device_check_calibration.png",
                "text": "Before starting the test, candidates undergo mandatory hardware verification: Webcam feed, Microphone audio, Browser resolution, and Network latency. The AI engine validates camera framing and illumination using a centering reticle. Once verified, the candidate clicks 'Enter Exam Portal'."
            },
            {
                "title": "Step 3.3: Candidate Exam Portal & Lockdown Interface",
                "image": "step_06_exam_portal_overview.png",
                "text": "The Exam Portal initiates browser lockdown enforcement (fullscreen monitoring, focus tracking). The interface displays the test title, live remaining timer, candidate live camera feed preview, and a numbered question navigation palette."
            },
            {
                "title": "Step 3.4: Answering Single-Choice Multiple Choice Questions (MCQ)",
                "image": "step_07_exam_mcq_single.png",
                "text": "For single-choice questions, candidates select their desired answer from the options. The response is immediately recorded and periodically autosaved. Candidates click 'Next' or use the palette to jump between questions."
            },
            {
                "title": "Step 3.5: Answering Multiple-Select Questions (MSQ)",
                "image": "step_08_exam_multi_correct.png",
                "text": "For questions with multiple valid solutions, candidates select all matching checkboxes. The system keeps track of all checked options and updates question palette status in real time."
            },
            {
                "title": "Step 3.6: Coding Environment with Monaco Editor & Real-Time Telemetry",
                "image": "step_09_exam_coding_monaco.png",
                "text": "For coding assessments, candidates write solutions within the embedded Monaco Editor (the editor powering VS Code) with syntax highlighting and auto-indentation. Keystroke timing, paste anomalies, and text deltas are continuously captured for academic integrity replay."
            },
            {
                "title": "Step 3.7: Answering Descriptive & Essay Questions",
                "image": "step_10_exam_descriptive.png",
                "text": "Descriptive questions provide a multi-line response area. Candidates write comprehensive theoretical explanations, which are continuously synchronized to the backend via fault-tolerant autosave."
            },
            {
                "title": "Step 3.8: Final Exam Submission Modal",
                "image": "step_11_exam_submit_confirmation.png",
                "text": "When finished, candidates click 'Submit Exam' in the header. A confirmation dialog appears summarizing the total answered questions and warning that further changes cannot be made after confirmation. The candidate clicks 'Confirm & Submit Exam'."
            },
            {
                "title": "Step 3.9: Submission Confirmation & Receipt",
                "image": "step_12_exam_submitted_receipt.png",
                "text": "Following submission, the candidate arrives at the completion receipt screen, which records the exam completion timestamp and guarantees that all responses and telemetry were successfully persisted."
            }
        ]
    },
    {
        "flow": "Flow 4: Administrator / Instructor Management & Invigilation Flow",
        "description": "This flow guides instructors through test authoring, question management, seating plan configuration, live proctoring supervision, and graded sheet reviews.",
        "steps": [
            {
                "title": "Step 4.1: Administrator Dashboard",
                "image": "step_14_admin_dashboard.png",
                "text": "The Administrator Dashboard provides an overview of active exams, total candidate registrations, and quick action shortcuts to the Test Studio and Live Monitor."
            },
            {
                "title": "Step 4.2: Instructor Test Studio Overview",
                "image": "step_15_admin_exams_studio.png",
                "text": "The Test Studio (/admin/exams) provides centralized management for all examinations. Instructors can create new exams, edit existing configurations, adjust question banks, and launch the invigilator monitor."
            },
            {
                "title": "Step 4.3: Exam Details & Metadata Configuration",
                "image": "step_15a_admin_exam_metadata.png",
                "text": "In the 'Metadata & Details' tab, instructors configure the Test ID, title, description, testing hall location, maximum candidate capacity, and idle auto-logout timeout."
            },
            {
                "title": "Step 4.4: Question Bank Builder & Multi-Format Authoring",
                "image": "step_15b_admin_exam_questions.png",
                "text": "In the 'Question Maker' tab, instructors create and edit questions across multiple formats (MCQ, Multi-Correct, Monaco Coding, Descriptive) with points, starter code, and answer keys."
            },
            {
                "title": "Step 4.5: Seating Plan & Classroom Coordinate Upload",
                "image": "step_15c_admin_exam_seating.png",
                "text": "In the 'Seating CSV' tab, instructors upload or define classroom desk coordinates (row and column mappings). Proctora uses these coordinates to deconflict question orders so adjacent students receive different question sequences."
            },
            {
                "title": "Step 4.6: Real-Time Live Invigilation & Proctoring Monitor",
                "image": "step_17_admin_live_monitor.png",
                "text": "The Live Invigilation Monitor (/admin/exams/:examId/monitor) connects to real-time WebSockets. Instructors monitor student status cards, AI-detected infraction flags (looking away, missing face, multiple faces), tab focus loss, and live cheating risk scores."
            },
            {
                "title": "Step 4.7: Candidate Graded Answer Sheet & Telemetry Inspector",
                "image": "step_18_admin_sheet_inspector.png",
                "text": "Instructors click 'Inspect Sheet' on any candidate card to open the Graded Sheet Viewer. This modal provides question-by-question candidate answers, code submissions, awarded points, and timestamped infraction logs."
            }
        ]
    }
]

def build_docx():
    print("Generating Refined Word Document: Proctora_User_Manual.docx...")
    doc = docx.Document()

    # Set 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Title
    title = doc.add_heading("Proctora - System User Manual", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub = doc.add_paragraph("Real-Time Online Examination System with Proctoring & Suspicious Event Detection")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(12)
    sub.runs[0].font.italic = True
    sub.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    # Authors Table Box
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, a in enumerate(authors):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{i+1}. {a['name']}\n")
        r1.bold = True
        r1.font.size = Pt(10.5)
        r2 = p.add_run(f"({a['roll']})")
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()
    inst = doc.add_paragraph()
    inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = inst.add_run("Department of Information Technology\nNational Institute of Technology Karnataka (NITK), Surathkal\nSoftware Engineering Course Project")
    r_inst.font.size = Pt(9.5)
    r_inst.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # Table of Contents
    doc.add_heading("Table of Contents / Overview of Flows", level=1)
    for sec in sections:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r_flow = p.add_run(sec["flow"])
        r_flow.bold = True
        for st in sec["steps"]:
            p2 = doc.add_paragraph(style='List Bullet 2')
            p2.paragraph_format.space_after = Pt(2)
            p2.add_run(st["title"])

    doc.add_page_break()

    # Iterate flows
    for s_idx, sec in enumerate(sections):
        h1 = doc.add_heading(sec["flow"], level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(4)

        p_desc = doc.add_paragraph(sec["description"])
        p_desc.paragraph_format.space_after = Pt(14)
        p_desc.runs[0].font.color.rgb = RGBColor(51, 65, 85)

        for st in sec["steps"]:
            h2 = doc.add_heading(st["title"], level=2)
            h2.paragraph_format.space_before = Pt(10)
            h2.paragraph_format.space_after = Pt(4)

            p_st = doc.add_paragraph(st["text"])
            p_st.paragraph_format.space_after = Pt(8)

            img_path = os.path.join(SCREENSHOTS_DIR, st["image"])
            if os.path.exists(img_path):
                try:
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_after = Pt(14)
                    run_img = p_img.add_run()
                    run_img.add_picture(img_path, width=Inches(5.8))
                except Exception as e:
                    print(f"Warning: could not add image {st['image']}: {e}")
            else:
                doc.add_paragraph(f"[Screenshot: {st['image']}]")

        if s_idx < len(sections) - 1:
            doc.add_page_break()

    doc.save(DOCX_OUTPUT)
    print(f"✅ Refined Word document generated at: {DOCX_OUTPUT}")

def build_html():
    print("Generating Refined HTML for PDF conversion: Proctora_User_Manual.html...")
    html_parts = []
    html_parts.append("""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Proctora - System User Manual</title>
<style>
  @page {
    size: A4;
    margin: 15mm 15mm 15mm 15mm;
  }
  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    line-height: 1.5;
    padding: 0;
    margin: 0 auto;
    max-width: 820px;
  }
  .header-card {
    text-align: center;
    border-bottom: 2px solid #e2e8f0;
    padding-bottom: 20px;
    margin-bottom: 25px;
  }
  h1 {
    font-size: 26px;
    margin-bottom: 4px;
    color: #0f172a;
    letter-spacing: -0.5px;
  }
  .subtitle {
    color: #475569;
    font-size: 13.5px;
    margin-bottom: 18px;
  }
  .authors-grid {
    display: flex;
    justify-content: center;
    gap: 20px;
    margin: 16px 0 12px;
  }
  .author-box {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 8px 14px;
    text-align: center;
  }
  .author-name {
    font-weight: 600;
    font-size: 13px;
    color: #0f172a;
  }
  .author-roll {
    font-size: 12px;
    color: #64748b;
  }
  .institution-note {
    font-size: 11.5px;
    color: #64748b;
    margin-top: 6px;
  }
  .toc-card {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    padding: 16px 20px;
    margin-bottom: 30px;
    border-radius: 6px;
    page-break-after: always;
  }
  .toc-card h3 {
    margin-top: 0;
    margin-bottom: 10px;
    font-size: 15px;
    color: #0f172a;
  }
  .toc-card ul {
    margin: 0;
    padding-left: 18px;
  }
  .toc-card li {
    margin-bottom: 4px;
    font-size: 12.5px;
  }
  .flow-section {
    margin-bottom: 30px;
    page-break-before: always;
  }
  .flow-section:first-of-type {
    page-break-before: auto;
  }
  h2 {
    font-size: 19px;
    color: #1e3a8a;
    border-bottom: 1.5px solid #cbd5e1;
    padding-bottom: 5px;
    margin-top: 15px;
    margin-bottom: 8px;
  }
  .flow-desc {
    font-size: 13px;
    color: #475569;
    margin-top: 0;
    margin-bottom: 18px;
  }
  .step-card {
    margin-bottom: 24px;
    page-break-inside: avoid;
    break-inside: avoid;
  }
  h3 {
    font-size: 14.5px;
    color: #0f172a;
    margin-top: 0;
    margin-bottom: 4px;
  }
  p.step-text {
    font-size: 12.5px;
    color: #334155;
    margin-top: 0;
    margin-bottom: 10px;
    line-height: 1.45;
  }
  .step-img {
    width: 100%;
    max-width: 790px;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
    display: block;
    margin: 0 auto 12px;
    box-shadow: 0 1px 3px rgba(0,0,0,0.06);
  }
</style>
</head>
<body>

<div class="header-card">
  <h1>Proctora - System User Manual</h1>
  <div class="subtitle">Real-Time Online Examination System with Proctoring & Suspicious Event Detection</div>

  <div class="authors-grid">
""")

    for a in authors:
        html_parts.append(f"""    <div class="author-box">
      <div class="author-name">{a['name']}</div>
      <div class="author-roll">{a['roll']}</div>
    </div>""")

    html_parts.append("""  </div>
  <div class="institution-note">
    Department of Information Technology &bull; National Institute of Technology Karnataka (NITK), Surathkal<br/>
    Software Engineering Course Project
  </div>
</div>

<div class="toc-card">
  <h3>Table of Contents / Overview of Flows</h3>
  <ul>
""")

    for sec in sections:
        html_parts.append(f"    <li><strong>{sec['flow']}</strong><ul>")
        for st in sec["steps"]:
            html_parts.append(f"      <li>{st['title']}</li>")
        html_parts.append("    </ul></li>")

    html_parts.append("""  </ul>
</div>
""")

    for sec in sections:
        html_parts.append(f'<div class="flow-section">')
        html_parts.append(f'  <h2>{sec["flow"]}</h2>')
        html_parts.append(f'  <p class="flow-desc">{sec["description"]}</p>')

        for st in sec["steps"]:
            img_rel = os.path.relpath(os.path.join(SCREENSHOTS_DIR, st["image"]), DOCS_DIR)
            html_parts.append(f'  <div class="step-card">')
            html_parts.append(f'    <h3>{st["title"]}</h3>')
            html_parts.append(f'    <p class="step-text">{st["text"]}</p>')
            html_parts.append(f'    <img class="step-img" src="{img_rel}" alt="{st["title"]}"/>')
            html_parts.append(f'  </div>')

        html_parts.append(f'</div>')

    html_parts.append("""
</body>
</html>
""")

    with open(HTML_OUTPUT, "w", encoding="utf-8") as f:
        f.write("\n".join(html_parts))
    print(f"✅ Refined HTML file generated at: {HTML_OUTPUT}")

build_docx()
build_html()
sys.exit(0)
